import csv
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
import pdf_tables_to_csv as m  # noqa: E402
import make_samples  # noqa: E402


@pytest.fixture(scope="module")
def pdfs(tmp_path_factory):
    d = tmp_path_factory.mktemp("pdfs")
    st, mx = d / "statement.pdf", d / "mixed.pdf"
    rows = make_samples.build_statement(st)
    make_samples.build_mixed(mx)
    return st, mx, rows


@pytest.mark.parametrize("spec,expected", [
    (None, [0, 1, 2]), ("1", [0]), ("2-3", [1, 2]), ("1,3", [0, 2]), ("2-", [1, 2]), ("3,1,1", [0, 2])])
def test_parse_pages(spec, expected):
    assert m.parse_pages(spec, 3) == expected


@pytest.mark.parametrize("spec", ["0", "4", "3-2", "1-9"])
def test_parse_pages_bad(spec):
    with pytest.raises(ValueError):
        m.parse_pages(spec, 3)


@pytest.mark.parametrize("text,expected", [
    ("£1,234.50", 1234.5), ("(12.00)", -12.0), ("50.00 DR", -50.0), ("50.00 CR", 50.0),
    ("-3.5", -3.5), ("120", 120), ("€9.99", 9.99), ("Card payment 4", "Card payment 4"),
    ("01/02/2026", "01/02/2026"), ("", ""), ("-", "-"), ("ABC123", "ABC123")])
def test_to_number(text, expected):
    assert m.to_number(text) == expected


def test_clean_cell():
    assert m.clean_cell("  Card\npayment   4 ") == "Card payment 4"
    assert m.clean_cell(None) == ""


def test_statement_found_on_both_pages(pdfs):
    st, _, _ = pdfs
    tables = m.extract_tables(st)
    assert [t["page"] for t in tables] == [1, 2]


def test_merge_keeps_one_header_and_every_row(pdfs):
    st, _, rows = pdfs
    merged = m.merge_tables(m.extract_tables(st))
    assert merged[0] == make_samples.HEADER
    assert merged[1:] == rows
    assert sum(1 for r in merged if r == make_samples.HEADER) == 1


def test_csv_output_roundtrip(pdfs, tmp_path):
    st, _, rows = pdfs
    [out] = m.run(st, tmp_path / "s")
    with open(out, encoding="utf-8-sig") as f:
        got = list(csv.reader(f))
    assert got[0] == make_samples.HEADER and len(got) == 61 and got[-1] == rows[-1]


def test_numbers_option(pdfs, tmp_path):
    _, mx, _ = pdfs
    merged = m.convert_rows(m.merge_tables(m.extract_tables(mx)), numbers=True)
    assert merged[0] == ["Product", "Unit price", "Qty"]
    assert merged[1] == ["Widget A", 1250.0, 2]
    assert merged[2][1] == -15.0 and merged[3][1] == -50.0
    # different table kept with its own header after a blank row
    i = merged.index([])
    assert merged[i + 1] == ["Region", "Customers"] and merged[i + 2] == ["North", 120]


def test_separate_files(pdfs, tmp_path):
    _, mx, _ = pdfs
    outs = m.run(mx, tmp_path / "mx", separate=True)
    assert len(outs) == 2 and all(p.exists() for p in outs)


def test_xlsx(pdfs, tmp_path):
    from openpyxl import load_workbook
    st, _, _ = pdfs
    [out] = m.run(st, tmp_path / "s", xlsx=True, numbers=True)
    ws = load_workbook(out)["tables"]
    assert ws["A1"].value == "Date" and ws.max_row == 61
    assert isinstance(ws["E2"].value, float) and ws.freeze_panes == "A2"


def test_page_filter(pdfs):
    st, _, _ = pdfs
    assert [t["page"] for t in m.extract_tables(st, pages="2")] == [2]


def test_no_tables_message(tmp_path):
    from reportlab.pdfgen import canvas
    p = tmp_path / "plain.pdf"
    c = canvas.Canvas(str(p)); c.drawString(100, 700, "Just text"); c.save()
    with pytest.raises(SystemExit, match="No tables found"):
        m.run(p)


def test_cli(pdfs, tmp_path, capsys):
    st, _, _ = pdfs
    assert m.main([str(st), "-o", str(tmp_path / "c")]) == 0
    assert (tmp_path / "c.csv").exists()
