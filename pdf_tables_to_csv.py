#!/usr/bin/env python3
"""pdf-tables-to-csv: pull every table out of a PDF into clean CSV (or one Excel file).

Built for the PDFs people actually get sent: bank statements, supplier price lists,
invoices and reports whose tables run across many pages with the header repeated.

    python pdf_tables_to_csv.py statement.pdf                # -> statement.csv
    python pdf_tables_to_csv.py statement.pdf --xlsx         # -> statement.xlsx
    python pdf_tables_to_csv.py report.pdf --separate        # one CSV per table
    python pdf_tables_to_csv.py report.pdf --pages 2-5 --numbers

MIT licence. Made by TCH Works (https://www.tchworks.co.uk).
"""
from __future__ import annotations

import argparse
import csv
import re
import sys
from pathlib import Path

import pdfplumber

__version__ = "1.0.0"

_MONEY = re.compile(r"^\(?-?[£$€]?\s?-?[\d,]*\.?\d+\)?(\s?(CR|DR|cr|dr))?$")


def parse_pages(spec: str | None, n_pages: int) -> list[int]:
    """'1,3-5' -> [0, 2, 3, 4] (zero-based). None -> all pages."""
    if not spec:
        return list(range(n_pages))
    out: list[int] = []
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            start, end = int(a), int(b) if b else n_pages
        else:
            start = end = int(part)
        if start < 1 or end > n_pages or start > end:
            raise ValueError(f"page range {part!r} is outside 1-{n_pages}")
        out.extend(range(start - 1, end))
    return sorted(set(out))


def clean_cell(value) -> str:
    """Collapse wrapped lines and stray spaces inside a cell."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def to_number(text: str):
    """'£1,234.50' -> 1234.5, '(12.00)' -> -12.0, '50.00 DR' -> -50.0.
    Returns the original text if it is not clearly a number."""
    t = text.strip()
    if not t or not _MONEY.match(t):
        return text
    negative = t.startswith("(") and t.endswith(")")
    suffix = re.search(r"(CR|DR)$", t, re.I)
    if suffix and suffix.group(1).upper() == "DR":
        negative = True
    digits = re.sub(r"[^\d.\-]", "", t)
    if digits.count("-") > 1 or digits in ("", "-", ".", "-."):
        return text
    if digits.startswith("-"):
        negative, digits = True, digits[1:]
    try:
        n = float(digits)
    except ValueError:
        return text
    n = -n if negative else n
    return int(n) if n.is_integer() and "." not in digits else n


def _is_blank(row: list[str]) -> bool:
    return all(c == "" for c in row)


def extract_tables(pdf_path: str | Path, pages: str | None = None,
                   text_strategy: bool = False) -> list[dict]:
    """Return [{'page': 1, 'rows': [[...], ...]}, ...] for every table found."""
    settings = None
    if text_strategy:
        settings = {"vertical_strategy": "text", "horizontal_strategy": "text"}
    found = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for i in parse_pages(pages, len(pdf.pages)):
            page = pdf.pages[i]
            tables = page.extract_tables(settings) if settings else page.extract_tables()
            for t in tables:
                rows = [[clean_cell(c) for c in r] for r in t]
                rows = [r for r in rows if not _is_blank(r)]
                if rows:
                    found.append({"page": i + 1, "rows": rows})
    return found


def merge_tables(tables: list[dict]) -> list[list[str]]:
    """Join tables that share a header (e.g. a statement split over pages).
    The repeated header is kept once; tables with a different header are
    appended after a blank row with their own header."""
    merged: list[list[str]] = []
    current_header = None
    for t in tables:
        rows = t["rows"]
        header, body = rows[0], rows[1:]
        if current_header is not None and header == current_header:
            merged.extend(body)
            continue
        if current_header is not None and len(header) == len(current_header) \
                and not _looks_like_header(header):
            # continuation page whose header was not repeated
            merged.extend(rows)
            continue
        if merged:
            merged.append([])
        merged.append(header)
        merged.extend(body)
        current_header = header
    return merged


def _looks_like_header(row: list[str]) -> bool:
    filled = [c for c in row if c]
    if not filled:
        return False
    numeric = sum(1 for c in filled if not isinstance(to_number(c), str))
    return numeric == 0


def convert_rows(rows: list[list[str]], numbers: bool) -> list[list]:
    if not numbers:
        return rows
    out = []
    for idx, r in enumerate(rows):
        if idx == 0 or _looks_like_header(r):
            out.append(r)
        else:
            out.append([to_number(c) for c in r])
    return out


def write_csv(rows: list[list], path: Path) -> None:
    # utf-8-sig so Excel opens £ signs correctly
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        csv.writer(f).writerows(rows)


def write_xlsx(sheets: list[tuple[str, list[list]]], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(title=name[:31])
        for r in rows:
            ws.append(r)
        if rows:
            for c in ws[1]:
                c.font = Font(bold=True)
            ws.freeze_panes = "A2"
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 60)
    wb.save(path)


def run(pdf: Path, out: Path | None = None, pages: str | None = None, separate: bool = False,
        xlsx: bool = False, numbers: bool = False, text_strategy: bool = False) -> list[Path]:
    tables = extract_tables(pdf, pages, text_strategy)
    if not tables:
        raise SystemExit(
            f"No tables found in {pdf.name}. If it is a scanned image it needs OCR first; "
            "if the table has no lines, try --text.")
    stem = out if out else pdf.with_suffix("")
    written: list[Path] = []
    if separate:
        groups = [(f"p{t['page']}_t{n + 1}", convert_rows(t["rows"], numbers))
                  for n, t in enumerate(tables)]
    else:
        groups = [("tables", convert_rows(merge_tables(tables), numbers))]
    if xlsx:
        p = Path(str(stem) + ".xlsx")
        write_xlsx(groups, p)
        written.append(p)
    else:
        for name, rows in groups:
            p = Path(str(stem) + (f"_{name}" if separate else "") + ".csv")
            write_csv(rows, p)
            written.append(p)
    return written


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Pull every table out of a PDF into CSV or Excel.")
    ap.add_argument("pdf", type=Path)
    ap.add_argument("-o", "--out", type=Path, help="output path without extension")
    ap.add_argument("--pages", help="e.g. 1,3-5 (default: all)")
    ap.add_argument("--separate", action="store_true", help="one file/sheet per table")
    ap.add_argument("--xlsx", action="store_true", help="write one Excel workbook")
    ap.add_argument("--numbers", action="store_true",
                    help="turn £1,234.50 / (12.00) / 50.00 DR into real numbers")
    ap.add_argument("--text", action="store_true", help="for tables without ruled lines")
    ap.add_argument("--version", action="version", version=__version__)
    a = ap.parse_args(argv)
    if not a.pdf.exists():
        ap.error(f"{a.pdf} not found")
    for p in run(a.pdf, a.out, a.pages, a.separate, a.xlsx, a.numbers, a.text):
        print(p)
    return 0


if __name__ == "__main__":
    sys.exit(main())
